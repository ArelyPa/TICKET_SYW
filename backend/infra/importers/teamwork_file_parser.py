"""Parser de Capa 2 (spec 041) para el reporte estándar de tareas de Teamwork en Excel/CSV.

Homologa cada fila a un `dict` crudo (sin resolver contra la base de datos) con las mismas
claves que emite `teamwork_api_client.py`, para que `teamwork_import_service.py` (Capa 1) sea
agnóstico al origen. Sin dependencias nuevas: `openpyxl` (ya aprobado, spec 034) para `.xlsx`,
módulo estándar `csv` para `.csv`.
"""
import csv
import io
from datetime import date, datetime
from typing import IO

import openpyxl

REQUIRED_COLUMNS = (
    "Company name", "Project", "Task list", "ID", "Task name", "Task description",
    "Start date", "Due date", "Assigned to", "Created by", "Time estimate", "Parent task ID",
)


class TeamworkFileParseError(Exception):
    def __init__(self, message: str, missing_columns: list[str] | None = None) -> None:
        super().__init__(message)
        self.missing_columns = missing_columns or []


def parse_file(filename: str, content: bytes) -> list[dict]:
    """Punto de entrada único: decide el parser según la extensión de `filename`."""
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        return _parse_csv(content)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise TeamworkFileParseError("Formato de archivo no soportado (use .xlsx o .csv)")


def _parse_xlsx(content: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise TeamworkFileParseError("El archivo no tiene filas")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    _validate_headers(headers)
    return [_row_to_dict(headers, row, i + 2) for i, row in enumerate(rows[1:])]


def _parse_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise TeamworkFileParseError("El archivo no tiene filas")
    headers = [h.strip() for h in rows[0]]
    _validate_headers(headers)
    return [_row_to_dict(headers, row, i + 2) for i, row in enumerate(rows[1:])]


def _validate_headers(headers: list[str]) -> None:
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise TeamworkFileParseError(
            f"Faltan columnas requeridas del reporte de Teamwork: {', '.join(missing)}",
            missing_columns=missing,
        )


def _row_to_dict(headers: list[str], row: tuple, row_number: int) -> dict:
    values = dict(zip(headers, row))

    def get(col: str):
        v = values.get(col)
        if isinstance(v, str):
            v = v.strip()
        return v if v not in ("", None) else None

    def get_date(col: str):
        v = values.get(col)
        if isinstance(v, (datetime, date)):
            return v.date() if isinstance(v, datetime) else v
        if isinstance(v, str) and v.strip():
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
                try:
                    return datetime.strptime(v.strip(), fmt).date()
                except ValueError:
                    continue
        return None

    def get_int(col: str, default: int = 0) -> int:
        v = values.get(col)
        try:
            return int(v) if v not in ("", None) else default
        except (TypeError, ValueError):
            return default

    external_id = get("ID")
    parent_id = get("Parent task ID")
    return {
        "row_number": row_number,
        "external_id": str(int(external_id)) if isinstance(external_id, float) else (
            str(external_id) if external_id is not None else None),
        "company_name": get("Company name"),
        "project_name": get("Project"),
        "list_name": get("Task list"),
        "title": get("Task name"),
        "description": get("Task description") or "",
        "start_date": get_date("Start date"),
        "due_date": get_date("Due date"),
        "assignee_name": get("Assigned to"),
        "requester_name": get("Created by"),
        "estimated_minutes": get_int("Time estimate", default=0),
        "parent_external_id": str(int(parent_id)) if isinstance(parent_id, float) else (
            str(parent_id) if parent_id is not None else None),
    }
