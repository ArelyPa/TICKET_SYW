"""Parser de Capa 2 (spec 042) para el reporte estándar de tiempos mensuales de Teamwork en
Excel/CSV. Homologa cada fila a un `dict` crudo (sin resolver contra la base de datos) para que
`time_import_service.py` (Capa 1) sea agnóstico al origen — mismo patrón que
`teamwork_file_parser.py` de spec 041. Sin dependencias nuevas: `openpyxl` (ya aprobado) para
`.xlsx`, módulo estándar `csv` para `.csv`.
"""
import csv
import io
from datetime import date, datetime
from typing import IO

import openpyxl

# Columnas siempre requeridas; además se exige al menos una de cada grupo alternativo
# (research.md/data-model.md): duración (`Hours`/`Minutes`/`Decimal hours`), Usuario
# (`Who` o `First name`+`Last name` o `User ID`) y Tarea (`Task` o `Task ID`).
_ALWAYS_REQUIRED = ("ID", "Date/time", "Company", "Project", "Description")
_DURATION_COLUMNS = ("Hours", "Minutes", "Decimal hours")
_WHO_COLUMN_GROUPS = (("Who",), ("First name", "Last name"), ("User ID",))
_TASK_COLUMNS = ("Task", "Task ID")


class TeamworkTimeFileParseError(Exception):
    def __init__(self, message: str, missing_columns: list[str] | None = None) -> None:
        super().__init__(message)
        self.missing_columns = missing_columns or []


def parse_file(filename: str, content: bytes) -> list[dict]:
    """Punto de entrada único: decide el parser según la extensión de `filename`."""
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        return _parse_csv(content)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise TeamworkTimeFileParseError("Formato de archivo no soportado (use .xlsx o .csv)")


def _parse_xlsx(content: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise TeamworkTimeFileParseError("El archivo no tiene filas")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    _validate_headers(headers)
    return [_row_to_dict(headers, row, i + 2) for i, row in enumerate(rows[1:])]


def _parse_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise TeamworkTimeFileParseError("El archivo no tiene filas")
    headers = [h.strip() for h in rows[0]]
    _validate_headers(headers)
    return [_row_to_dict(headers, row, i + 2) for i, row in enumerate(rows[1:])]


def _validate_headers(headers: list[str]) -> None:
    missing = [c for c in _ALWAYS_REQUIRED if c not in headers]
    if not any(c in headers for c in _DURATION_COLUMNS):
        missing.append("Hours|Minutes|Decimal hours")
    if not any(all(c in headers for c in group) for group in _WHO_COLUMN_GROUPS):
        missing.append("Who|First name+Last name|User ID")
    if not any(c in headers for c in _TASK_COLUMNS):
        missing.append("Task|Task ID")
    if missing:
        raise TeamworkTimeFileParseError(
            f"Columnas requeridas faltantes: {', '.join(missing)}", missing_columns=missing)


def _row_to_dict(headers: list[str], row: tuple, row_number: int) -> dict:
    values = dict(zip(headers, row))

    who_name = _clean(values.get("Who"))
    if not who_name:
        first, last = _clean(values.get("First name")), _clean(values.get("Last name"))
        if first or last:
            who_name = " ".join(p for p in (first, last) if p)

    task_reference = _clean(values.get("Task ID")) or _clean(values.get("Task"))

    return {
        "row_number": row_number,
        "external_time_id": _clean(values.get("ID")),
        "started_at": _parse_datetime(values.get("Date/time")),
        "ended_at": _parse_datetime(values.get("End date/time")),
        "duration_minutes": _resolve_duration_minutes(values),
        "who_name": who_name,
        "who_external_id": _clean(values.get("User ID")),
        "company_name": _clean(values.get("Company")),
        "project_name": _clean(values.get("Project")),
        "task_reference": task_reference,
        "description": _clean(values.get("Description")) or "",
    }


def _resolve_duration_minutes(values: dict) -> int | None:
    """research.md Decisión 5: `Decimal hours` > `Hours`+`Minutes` > sin dato (conflicto,
    nunca se adivina sumando/promediando las tres)."""
    decimal_hours = _to_float(values.get("Decimal hours"))
    if decimal_hours is not None:
        return round(decimal_hours * 60)
    hours = _to_float(values.get("Hours"))
    minutes = _to_float(values.get("Minutes"))
    if hours is not None or minutes is not None:
        return round((hours or 0) * 60 + (minutes or 0))
    return None


def _to_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_datetime(value) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None
